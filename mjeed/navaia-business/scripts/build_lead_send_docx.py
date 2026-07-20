#!/usr/bin/env python3
"""Build a one-lead manual-send DOCX for a real CRM lead: every token already filled.

Where build_manual_send_docx.py prints the blank templates, this prints the finished
messages for ONE lead, so the operator can copy them straight out on a phone.

Touch 1 comes from outreach.render_lead() — the same function the automated pipeline
calls — so what is printed here is byte-exact with what an automated send would deliver.
Touches 2 and 3 have no pipeline equivalent (the pipeline only ever sends Touch 1), so
they are filled here using the same public helpers render_lead uses: short_name(),
lina_compose.compose_block() and outreach.EMAIL_PAIN.

Reads the CRM. Sends nothing, spends nothing, writes nothing back.

    python scripts/build_lead_send_docx.py --vertical "Training Institutes"
    python scripts/build_lead_send_docx.py --vertical "Training Institutes" --match الآن
    python scripts/build_lead_send_docx.py --vertical "Training Institutes" --list
"""
from __future__ import annotations

import argparse
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
if (sys.stdout.encoding or "").lower().replace("-", "") != "utf8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Pt

import lina_compose
import outreach
import pipeline_prep as prep
from build_manual_send_docx import (EMAIL_STYLE, arabic_block, heading, kv_table,
                                    note, parse_templates_md)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Vertical -> the heading it appears under in 04_outreach_templates.md, matched loosely so
# a renumbered or retitled section still resolves.
_SECTION_HINT = {
    "Training Institutes": "Training Institutes",
    "Real Estate": "Real Estate",
    "Contracting & Facilities": "Contracting",
    "Private Clinics": "Private Specialty Clinics",
    "Finance & Debt Collection": "Finance",
}


def _all_leads(vertical: str) -> list[dict]:
    """Every one of Mjeed's leads in this vertical, whatever their leadStatus.

    Mirrors pipeline_prep.not_contacted_leads() field for field, minus the status filter:
    a sheet is often wanted for a company that was ALREADY contacted (to re-read what went
    out, or to send a follow-up), and restricting to the outreach queue made naming such a
    company match nothing. Same reasoning as the --only fix in enrich_company_size.py.
    """
    out = []
    for p in prep._crm_people():
        if p.get("sector") != vertical:
            continue
        name = " ".join(x for x in [(p.get("name") or {}).get("firstName"),
                                    (p.get("name") or {}).get("lastName")] if x).strip()
        company = (p.get("company") or {}).get("name") or name
        phone = (p.get("phones") or {}).get("primaryPhoneNumber") or ""
        out.append({
            "person_id": p["id"],
            "company": company,
            "contact_name": name if name != company else "",
            "sector": p["sector"],
            "email": (p.get("emails") or {}).get("primaryEmail") or "",
            "phone": phone,
            "website": ((p.get("company") or {}).get("domainName") or {}).get("primaryLinkUrl") or "",
            "pain_line": prep.scrape_pain_index().get(prep._phone9(phone), ""),
            "lead_status": p.get("leadStatus") or "",
        })
    return out


def pick_lead(vertical: str, match: str | None, show_list: bool) -> dict | None:
    leads = _all_leads(vertical)
    if not leads:
        raise SystemExit(f"No leads in {vertical}.")
    if show_list:
        for i, l in enumerate(leads, 1):
            print(f"{i:>3}. {l['company']}  |  {l.get('email') or '-'}  |  "
                  f"{l.get('phone') or '-'}  |  {l.get('lead_status') or '-'}")
        return None
    if match:
        hits = [l for l in leads if match in l["company"]]
        if not hits:
            raise SystemExit(f"No {vertical} lead matches {match!r}. Use --list to see them.")
        return hits[0]
    # Default: the first lead that can actually receive both channels.
    both = [l for l in leads if l.get("email") and l.get("phone")]
    return (both or leads)[0]


def other_touches(vertical: str, lead: dict) -> list[dict]:
    """Fill Touch 2 and 3 for this lead with the same substitutions render_lead makes."""
    verticals, _ = parse_templates_md()
    hint = _SECTION_HINT.get(vertical, vertical)
    section = next((v for k, v in verticals.items() if hint in k), None)
    if section is None:
        return []

    company = lead["company"]
    honorific = (f"الأستاذ {lead['contact_name']}" if lead.get("contact_name")
                 else f"القائمون على {outreach.short_name(company)} الكرام")
    _, meta = lina_compose.compose_block(
        outreach.VKEY[vertical], lead.get("pain_line") or lead.get("pain_hints") or "",
        use_llm=False, key=None)
    pain_phrase = meta["pain"] if meta.get("category") else outreach.EMAIL_PAIN[vertical]
    # Touch 1 leads with the company's own published facts when the sizing crawl found any
    # (outreach.site_opener), else the generic vertical line. Touches 2 and 3 must resolve
    # {trigger_line} the SAME way or one sheet opens on "بتغطيتكم 13 منطقةً…" and its own
    # follow-up opens on the generic sentence, which reads as two different senders.
    opener, _ = outreach.site_opener(lead)
    fills = {
        "{honorific+name}": honorific,
        "{inbound_context}": "",
        "{trigger_line}": opener or outreach.TRIGGER[vertical],
        "{lina_pain}": pain_phrase,
        "{اسم الشركة}": company, "{اسم العيادة}": company, "{اسم المعهد}": company,
    }
    out = []
    for t in section["touches"]:
        if t["label"].startswith("Touch 1"):
            continue
        body = t["body"]
        for k, v in fills.items():
            body = body.replace(k, v)
        out.append({**t, "body": body})
    return out


def build(lead: dict, vertical: str, r: dict, path: str) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(lead["company"], 0)
    note(doc, f"Manual-send sheet — {vertical}. Every token is already filled: copy the "
              f"blocks below exactly as they are.", italic=False)
    note(doc, "Touch 1 is rendered by the same code path as an automated send. "
              "Regenerate: python scripts/build_lead_send_docx.py "
              f'--vertical "{vertical}" --match {outreach.short_name(lead["company"])[:12]}')

    heading(doc, "Lead", 1)
    kv_table(doc, [
        ("Company", lead["company"]),
        ("CRM person_id", lead.get("person_id", "-")),
        ("Contact name", lead.get("contact_name") or "— (none in CRM; the greeting falls "
                                                     "back to القائمون على … الكرام)"),
        ("Email", lead.get("email") or "— none, so email cannot be sent"),
        ("WhatsApp number", r["phone_intl"] or "— none"),
        ("Lead status", lead.get("lead_status") or "-"),
        ("Opener", "the company's own published facts"
                   if r.get("opener_source", "generic") != "generic"
                   else "generic vertical line (no facts found on their site)"),
        ("Pain source", r["pain_source"]),
    ])
    if r["pain_source"] == "general":
        note(doc, "Pain source is 'general': no lead-specific pain matched, so the copy uses "
                  "the vertical's default pain. That is the same thing an automated send "
                  "would do — it is not a gap you need to fill in by hand.")

    # ---- email
    heading(doc, "Email", 1)
    if not r["email_subject"]:
        note(doc, "No email address on this lead — WhatsApp only.")
    else:
        kv_table(doc, [
            ("To", lead["email"]),
            ("From", "ops@navaia.sa"),
            ("Format", f'HTML, {EMAIL_STYLE["font"]} {EMAIL_STYLE["size_pt"]}pt, '
                       f'line spacing {EMAIL_STYLE["line_spacing"]}, RTL right-aligned'),
            ("Signature", "do not type one — appended server-side"),
            ("Send window", "Sun–Thu, ~10:00–12:00 AST"),
        ])
        heading(doc, "Touch 1 — day 0", 2)
        p = doc.add_paragraph()
        p.add_run("Subject: ").bold = True
        sr = p.add_run(r["email_subject"])
        sr.font.name = EMAIL_STYLE["font"]
        from build_manual_send_docx import _rtl
        _rtl(p, sr)
        # r["email_body"] is the HTML the pipeline sends; show the copy as it will read.
        plain = re.sub(r"<br\s*/?>", "\n", r["email_body"])
        plain = re.sub(r"<[^>]+>", "", plain).strip()
        arabic_block(doc, plain)

        for t in other_touches(vertical, lead):
            doc.add_paragraph()
            heading(doc, t["label"], 2)
            if t["subject"]:
                p = doc.add_paragraph()
                p.add_run("Subject: ").bold = True
                sr = p.add_run(t["subject"])
                sr.font.name = EMAIL_STYLE["font"]
                _rtl(p, sr)
            arabic_block(doc, t["body"])

    # ---- whatsapp
    doc.add_page_break()
    heading(doc, "WhatsApp — Touch 1", 1)
    kv_table(doc, [
        ("To", r["phone_intl"] or "— none"),
        ("Template", r["wa_template"]),
        ("Language", "ar"),
        ("Formatting", "none — WhatsApp renders the approved body as-is"),
    ])
    note(doc, "This is the message as it will arrive. If you send through the Business API "
              "you pass the five variables below instead; the fixed text comes from the "
              "approved template.")
    arabic_block(doc, r["wa_preview"])

    doc.add_paragraph()
    heading(doc, "The five variables", 2)
    kv_table(doc, [(f"{{{{{i}}}}}", v) for i, v in enumerate(r["wa_vars"], 1)])

    doc.save(path)
    print(f"wrote {path}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--vertical", default="Training Institutes")
    ap.add_argument("--match", help="substring of the company name")
    ap.add_argument("--list", action="store_true", help="list the leads and exit")
    args = ap.parse_args()

    if args.vertical not in outreach.VKEY:
        raise SystemExit(f"{args.vertical!r} is not an active vertical. "
                         f"Active: {', '.join(outreach.VKEY)}")

    lead = pick_lead(args.vertical, args.match, args.list)
    if lead is None:
        return
    r = outreach.render_lead(lead, use_llm=False, or_key=None)
    slug = re.sub(r"[^\w]+", "_", outreach.short_name(lead["company"]))[:40].strip("_")
    build(lead, args.vertical, r, os.path.join(ROOT, f"send_{slug}.docx"))


if __name__ == "__main__":
    main()
